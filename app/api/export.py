"""Export API endpoints for downloading content in various formats."""
import json
import io
import re
import zipfile
from datetime import datetime
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse, Response

from app.database import get_db
from app.db_utils import fetchone, fetchall
from app.api.auth import get_current_user_id

router = APIRouter()


def sanitize_filename(filename: str) -> str:
    """
    Sanitize a filename to prevent path traversal and invalid characters.

    - Removes path separators (/, \\)
    - Removes null bytes
    - Removes control characters
    - Replaces problematic characters with underscores
    - Limits length to 200 characters
    - Ensures the filename is not empty
    """
    if not filename:
        return "export"

    # Remove path components (prevent path traversal)
    filename = filename.replace("/", "_").replace("\\", "_")

    # Remove null bytes and control characters
    filename = re.sub(r'[\x00-\x1f\x7f]', '', filename)

    # Remove/replace characters that are problematic in filenames or headers
    # Keep alphanumeric, spaces, hyphens, underscores, and dots
    filename = re.sub(r'[^\w\s\-.]', '_', filename)

    # Collapse multiple underscores/spaces
    filename = re.sub(r'[_\s]+', '_', filename)

    # Remove leading/trailing underscores and spaces
    filename = filename.strip('_').strip()

    # Limit length (reserve space for extension)
    if len(filename) > 200:
        filename = filename[:200]

    # Ensure we have a valid filename
    return filename if filename else "export"


async def get_job_with_outputs(job_id: str, user_id: int) -> tuple[dict, list[dict], list[dict]]:
    """
    Fetch job data with outputs and stills.

    Returns (job, outputs, stills) tuple.
    """
    async with get_db() as db:
        # Get job details
        job = await fetchone(
            db,
            """
            SELECT id, original_filename, target_persona, asset_types,
                   cost_incurred, created_at, completed_at
            FROM jobs WHERE id = ? AND user_id = ?
            """,
            (job_id, user_id)
        )

        if not job:
            raise HTTPException(status_code=404, detail="Job not found")

        # Get outputs
        output_rows = await fetchall(
            db,
            """
            SELECT id, content_type, variation_number,
                   step1_draft, step2_edited, step3_final,
                   atoms_used, citations, warnings, quality_scores,
                   hook_variations
            FROM outputs WHERE job_id = ?
            ORDER BY content_type, variation_number
            """,
            (job_id,)
        )

        outputs = []
        for row in output_rows:
            outputs.append({
                "id": row["id"],
                "content_type": row["content_type"],
                "variation_number": row["variation_number"],
                "step1_draft": row["step1_draft"],
                "step2_edited": row["step2_edited"],
                "step3_final": row["step3_final"],
                "atoms_used": json.loads(row["atoms_used"]) if row["atoms_used"] else [],
                "citations": json.loads(row["citations"]) if row["citations"] else [],
                "warnings": json.loads(row["warnings"]) if row["warnings"] else [],
                "quality_scores": json.loads(row["quality_scores"]) if row["quality_scores"] else {},
                "hook_variations": json.loads(row["hook_variations"]) if row["hook_variations"] else [],
            })

        # Get stills
        still_rows = await fetchall(
            db,
            """
            SELECT id, still_type, content, source_location,
                   tags, persona_relevance, quote_attribution
            FROM stills WHERE job_id = ?
            ORDER BY still_type
            """,
            (job_id,)
        )

        stills = []
        for row in still_rows:
            stills.append({
                "id": row["id"],
                "type": row["still_type"],
                "content": row["content"],
                "source_location": row["source_location"],
                "tags": json.loads(row["tags"]) if row["tags"] else [],
                "persona_relevance": json.loads(row["persona_relevance"]) if row["persona_relevance"] else {},
                "quote_attribution": row["quote_attribution"],
            })

        return dict(job), outputs, stills


def format_markdown(job: dict, outputs: list[dict], atoms: list[dict]) -> str:
    """Format job results as markdown."""
    lines = []

    # Header
    lines.append(f"# {job['original_filename']}")
    lines.append("")
    lines.append(f"**Generated:** {job['created_at']}")
    lines.append(f"**Persona:** {job['target_persona']}")
    lines.append(f"**Cost:** ${job['cost_incurred']:.4f}")
    lines.append("")
    lines.append("---")
    lines.append("")

    # Group outputs by type
    output_groups = {}
    for output in outputs:
        content_type = output["content_type"]
        if content_type not in output_groups:
            output_groups[content_type] = []
        output_groups[content_type].append(output)

    # Render each content type
    for content_type, items in output_groups.items():
        type_title = content_type.replace("_", " ").title()
        lines.append(f"## {type_title}")
        lines.append("")

        for item in items:
            content = item.get("step3_final") or item.get("step2_edited") or item.get("step1_draft") or ""
            variation = item.get("variation_number", 1)
            scores = item.get("quality_scores", {})

            if content_type == "email_sequence":
                # Special formatting for email sequences
                lines.append(f"### Email {variation}")
                if item.get("subject"):
                    lines.append(f"**Subject:** {item.get('subject')}")
                if item.get("day") is not None:
                    lines.append(f"**Day:** {item.get('day')}")
                if item.get("purpose"):
                    lines.append(f"**Purpose:** {item.get('purpose')}")
                lines.append("")
                lines.append(content)
            else:
                lines.append(f"### Variation {variation}")
                if scores.get("overall_score"):
                    lines.append(f"*Quality Score: {scores['overall_score']}/100*")
                lines.append("")
                lines.append(content)

            # Hook variations if present
            if item.get("hook_variations"):
                lines.append("")
                lines.append("#### Hook Variations")
                for hook in item["hook_variations"]:
                    hook_type = hook.get("hook_type", "unknown").replace("_", " ").title()
                    lines.append(f"- **{hook_type}:** {hook.get('hook_text', '')}")

            # Warnings if present
            if item.get("warnings"):
                lines.append("")
                lines.append("**Warnings:**")
                for warning in item["warnings"]:
                    lines.append(f"- {warning}")

            lines.append("")
            lines.append("---")
            lines.append("")

    # Atoms section
    if atoms:
        lines.append("## Extracted Atoms")
        lines.append("")

        atom_groups = {}
        for atom in atoms:
            atom_type = atom["type"]
            if atom_type not in atom_groups:
                atom_groups[atom_type] = []
            atom_groups[atom_type].append(atom)

        for atom_type, items in atom_groups.items():
            lines.append(f"### {atom_type.title()}")
            lines.append("")
            for atom in items:
                content = atom["content"]
                if atom.get("quote_attribution"):
                    lines.append(f"> \"{content}\" — {atom['quote_attribution']}")
                else:
                    lines.append(f"- {content}")
            lines.append("")

    return "\n".join(lines)


def serialize_datetime(obj):
    """JSON serializer for datetime objects."""
    if hasattr(obj, 'isoformat'):
        return obj.isoformat()
    raise TypeError(f"Object of type {type(obj).__name__} is not JSON serializable")


def format_json(job: dict, outputs: list[dict], atoms: list[dict]) -> str:
    """Format job results as JSON."""
    return json.dumps({
        "job": {
            "id": job["id"],
            "original_filename": job["original_filename"],
            "target_persona": job["target_persona"],
            "cost_incurred": job["cost_incurred"],
            "created_at": job["created_at"],
            "completed_at": job["completed_at"],
        },
        "outputs": outputs,
        "atoms": atoms,
    }, indent=2, default=serialize_datetime)


def create_docx(job: dict, outputs: list[dict], atoms: list[dict]) -> bytes:
    """
    Create a DOCX file from job results.

    Uses python-docx if available, otherwise falls back to a simple format.
    """
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.style import WD_STYLE_TYPE

        doc = Document()

        # Title
        doc.add_heading(job["original_filename"], 0)

        # Metadata
        p = doc.add_paragraph()
        p.add_run("Generated: ").bold = True
        p.add_run(str(job["created_at"]))
        p = doc.add_paragraph()
        p.add_run("Persona: ").bold = True
        p.add_run(job["target_persona"])
        p = doc.add_paragraph()
        p.add_run("Cost: ").bold = True
        p.add_run(f"${job['cost_incurred']:.4f}")

        doc.add_paragraph()

        # Group outputs by type
        output_groups = {}
        for output in outputs:
            content_type = output["content_type"]
            if content_type not in output_groups:
                output_groups[content_type] = []
            output_groups[content_type].append(output)

        # Render each content type
        for content_type, items in output_groups.items():
            type_title = content_type.replace("_", " ").title()
            doc.add_heading(type_title, 1)

            for item in items:
                content = item.get("step3_final") or item.get("step2_edited") or item.get("step1_draft") or ""
                variation = item.get("variation_number", 1)
                scores = item.get("quality_scores", {})

                doc.add_heading(f"Variation {variation}", 2)

                if scores.get("overall_score"):
                    p = doc.add_paragraph()
                    p.add_run(f"Quality Score: {scores['overall_score']}/100").italic = True

                doc.add_paragraph(content)

                # Hook variations
                if item.get("hook_variations"):
                    doc.add_heading("Hook Variations", 3)
                    for hook in item["hook_variations"]:
                        hook_type = hook.get("hook_type", "unknown").replace("_", " ").title()
                        p = doc.add_paragraph()
                        p.add_run(f"{hook_type}: ").bold = True
                        p.add_run(hook.get("hook_text", ""))

                # Warnings
                if item.get("warnings"):
                    p = doc.add_paragraph()
                    p.add_run("Warnings:").bold = True
                    for warning in item["warnings"]:
                        doc.add_paragraph(f"  - {warning}")

        # Atoms section
        if atoms:
            doc.add_heading("Extracted Atoms", 1)

            atom_groups = {}
            for atom in atoms:
                atom_type = atom["type"]
                if atom_type not in atom_groups:
                    atom_groups[atom_type] = []
                atom_groups[atom_type].append(atom)

            for atom_type, items in atom_groups.items():
                doc.add_heading(atom_type.title(), 2)
                for atom in items:
                    content = atom["content"]
                    if atom.get("quote_attribution"):
                        p = doc.add_paragraph()
                        p.add_run(f'"{content}"').italic = True
                        p.add_run(f" — {atom['quote_attribution']}")
                    else:
                        doc.add_paragraph(f"  - {content}")

        # Save to bytes
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    except ImportError:
        # Fallback if python-docx not installed
        # Return the markdown as plain text
        content = format_markdown(job, outputs, atoms)
        return content.encode("utf-8")


@router.get("/export/{job_id}/markdown")
async def export_markdown(
    job_id: str,
    user_id: int = Depends(get_current_user_id),
):
    """Export job results as markdown file."""
    job, outputs, atoms = await get_job_with_outputs(job_id, user_id)

    content = format_markdown(job, outputs, atoms)
    base_name = sanitize_filename(job['original_filename'].rsplit('.', 1)[0])
    filename = f"{base_name}_export.md"

    return Response(
        content=content,
        media_type="text/markdown",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        }
    )


@router.get("/export/{job_id}/json")
async def export_json(
    job_id: str,
    user_id: int = Depends(get_current_user_id),
):
    """Export job results as JSON file."""
    job, outputs, atoms = await get_job_with_outputs(job_id, user_id)

    content = format_json(job, outputs, atoms)
    base_name = sanitize_filename(job['original_filename'].rsplit('.', 1)[0])
    filename = f"{base_name}_export.json"

    return Response(
        content=content,
        media_type="application/json",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        }
    )


@router.get("/export/{job_id}/docx")
async def export_docx(
    job_id: str,
    user_id: int = Depends(get_current_user_id),
):
    """Export job results as Word document."""
    job, outputs, atoms = await get_job_with_outputs(job_id, user_id)

    content = create_docx(job, outputs, atoms)
    base_name = sanitize_filename(job['original_filename'].rsplit('.', 1)[0])
    filename = f"{base_name}_export.docx"

    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        }
    )


@router.get("/export/{job_id}/zip")
async def export_zip(
    job_id: str,
    user_id: int = Depends(get_current_user_id),
):
    """Export all job outputs as a ZIP file containing markdown, JSON, and individual content files."""
    job, outputs, atoms = await get_job_with_outputs(job_id, user_id)

    base_name = sanitize_filename(job["original_filename"].rsplit(".", 1)[0])

    # Create ZIP in memory
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        # Add full markdown export
        zf.writestr(f"{base_name}_full_export.md", format_markdown(job, outputs, atoms))

        # Add JSON export
        zf.writestr(f"{base_name}_export.json", format_json(job, outputs, atoms))

        # Add individual content files
        for output in outputs:
            content_type = output["content_type"]
            variation = output.get("variation_number", 1)
            content = output.get("step3_final") or output.get("step2_edited") or output.get("step1_draft") or ""

            if content:
                filename = f"{content_type}_{variation}.txt"
                zf.writestr(f"content/{filename}", content)

                # If there are hook variations, export those too
                if output.get("hook_variations"):
                    hooks_md = f"# Hook Variations for {content_type.title()} {variation}\n\n"
                    for hook in output["hook_variations"]:
                        hook_type = hook.get("hook_type", "unknown").replace("_", " ").title()
                        hooks_md += f"## {hook_type}\n\n"
                        hooks_md += f"**Hook:** {hook.get('hook_text', '')}\n\n"
                        hooks_md += f"**Full Post:**\n{hook.get('full_post', '')}\n\n"
                        hooks_md += f"*{hook.get('explanation', '')}*\n\n"
                        hooks_md += "---\n\n"
                    zf.writestr(f"hooks/{content_type}_{variation}_hooks.md", hooks_md)

        # Add atoms export
        if atoms:
            atoms_md = "# Extracted Content Atoms\n\n"
            for atom in atoms:
                atoms_md += f"## {atom['type'].title()}\n\n"
                if atom.get("quote_attribution"):
                    atoms_md += f'> "{atom["content"]}" — {atom["quote_attribution"]}\n\n'
                else:
                    atoms_md += f"- {atom['content']}\n\n"
            zf.writestr("atoms.md", atoms_md)

    buffer.seek(0)
    filename = f"{base_name}_export.zip"

    return StreamingResponse(
        buffer,
        media_type="application/zip",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        }
    )
